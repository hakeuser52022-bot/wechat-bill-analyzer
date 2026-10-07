"""从个人简介 Markdown 生成 Word；默认拒绝覆盖。"""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--out',default=str(ROOT/'docs/个人简介（待填写）.docx'));ap.add_argument('--force',action='store_true');a=ap.parse_args()
    out=Path(a.out)
    if out.exists() and not a.force: raise SystemExit('文件已存在；请直接编辑，或显式使用 --force 覆盖。')
    d=Document();normal=d.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6)
    for line in (ROOT/'docs/intro.md').read_text(encoding='utf-8').splitlines():
        if not line.strip():continue
        if line.startswith('# '):p=d.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):p=d.add_paragraph(line[3:],style='Heading 2')
        else:p=d.add_paragraph(line)
        for r in p.runs:
            r.font.name='Microsoft YaHei';r.font.color.rgb=__import__('docx.shared',fromlist=['RGBColor']).RGBColor(0,0,0)
            rp=r._element.get_or_add_rPr();fonts=rp.find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');rp.append(fonts)
            fonts.set(qn('w:eastAsia'),'微软雅黑')
    d.core_properties.author='';d.core_properties.last_modified_by='';d.save(out)
if __name__=='__main__':main()
