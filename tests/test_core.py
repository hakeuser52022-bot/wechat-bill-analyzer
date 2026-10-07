# -*- coding: utf-8 -*-
"""
核心逻辑的自动化测试。

这些测试覆盖了开发过程中踩过的每一个坑，防止回归。
运行方式：

    # 方式一：用 pytest（推荐）
    pip install pytest
    python -m pytest tests/ -v

    # 方式二：不装任何测试框架，直接跑（内置极简 runner）
    python tests/test_core.py
"""
import os
import re
import sys
import json
sys.stdout.reconfigure(encoding="utf-8",errors="replace")
import tempfile
import collections

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..'))
sys.path.insert(0, os.path.join(ROOT, 'src'))

import analyze      # noqa: E402
import timeseg      # noqa: E402
import meals        # noqa: E402
import budget       # noqa: E402
from openpyxl import Workbook, load_workbook   # noqa: E402


# ---------------------------------------------------------------------------
# 测试夹具：构造一个最小但包含全部"坑"的账单
# ---------------------------------------------------------------------------
HEADERS = ['交易时间', '交易类型', '交易对方', '商品', '收/支', '金额(元)',
           '支付方式', '当前状态', '交易单号', '商户单号', '备注']


def _make_bill(path, preamble_lines=13):
    """写一个含抬头的账单：表头故意不放在第 1 行。"""
    from datetime import datetime
    rows = [
        # 正常食堂消费
        (datetime(2026, 9, 5, 8, 10), '商户消费', '示例大学', '示例食堂A',
         '支出', 8.50, '零钱', '支付成功'),
        (datetime(2026, 9, 5, 12, 30), '商户消费', '示例大学', '示例食堂B',
         '支出', 15.00, '零钱', '支付成功'),
        # 已全额退款：金额仍为正数、状态含"退款"，必须被剔除
        (datetime(2026, 9, 6, 21, 39), '商户消费', '淘宝平台商户', '商户单号123',
         '支出', 82.90, '零钱', '已全额退款'),
        # 部分退款
        (datetime(2026, 9, 7, 17, 45), '商户消费', '美团', '喜茶-美团App',
         '支出', 39.00, '零钱', '已退款(¥20.00)'),
        # 扫二维码付款给商家（形似转账，实为运动消费）
        (datetime(2026, 9, 8, 12, 13), '扫二维码付款', '示例羽毛球馆',
         '收款方备注:二维码收款', '支出', 6.00, '/', '已转账'),
        # 真正的转账给个人
        (datetime(2026, 9, 9, 14, 0), '转账', '演示同学A', '转账备注:微信转账',
         '支出', 30.00, '/', '已转账'),
        # 网约车
        (datetime(2026, 9, 10, 15, 0), '商户消费', '如祺出行', '先乘车后付款',
         '支出', 25.60, '零钱', '支付成功'),
        # 校园卡充值：预付性支出
        (datetime(2026, 9, 4, 12, 19), '商户消费', '示例大学信息中心', '校园卡充值',
         '支出', 200.00, '零钱', '支付成功'),
        # 亲人转账收入
        (datetime(2026, 9, 5, 12, 30), '转账', '母亲', '转账备注:微信转账',
         '收入', 1000.00, '/', '已存入零钱'),
        # 不计收支
        (datetime(2026, 9, 3, 16, 54), '零钱提现', '演示银行账户', '/',
         '/', 20.00, '/', '提现已到账'),
    ]
    wb = Workbook()
    ws = wb.active
    for i in range(1, preamble_lines + 1):
        ws.cell(row=i, column=1, value=f'账单抬头第 {i} 行')
    hr = preamble_lines + 1
    for j, h in enumerate(HEADERS, 1):
        ws.cell(row=hr, column=j, value=h)
    for i, r in enumerate(rows, hr + 1):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
    wb.save(path)
    return path


def _load():
    """加载夹具账单，返回 (rows, headers, header_row)。"""
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    try:
        _make_bill(path)
        return analyze.load_bill(path)
    finally:
        os.remove(path)


# ---------------------------------------------------------------------------
# 测试 1：表头自动定位
# ---------------------------------------------------------------------------
def test_find_header_row_not_first():
    """表头不在第 1 行时，必须能按内容找到（这是最容易被写死的假设）。"""
    rows, headers, hr = _load()
    assert hr == 14, f'表头应在第 14 行，实际 {hr}'
    assert '交易时间' in headers and '金额(元)' in headers
    assert len(rows) == 10, f'应读到 10 条记录，实际 {len(rows)}'


def test_find_header_row_varies():
    """抬头行数变化时同样要能找到（模拟不同账号导出）。"""
    for n in (5, 13, 20):
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        try:
            _make_bill(path, preamble_lines=n)
            rows, headers, hr = analyze.load_bill(path)
            assert hr == n + 1, f'抬头 {n} 行时表头应在第 {n+1} 行，实际 {hr}'
            assert len(rows) == 10
        finally:
            os.remove(path)


# ---------------------------------------------------------------------------
# 测试 2：退款剔除（问题 1）
# ---------------------------------------------------------------------------
def test_refund_excluded_from_net():
    """退款行必须从净支出中剔除，但保留在"流水"里。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)

    # 支出流水 = 8.50+15.00+82.90+39.00+6.00+30.00+25.60+200.00 = 407.00
    assert abs(res['gross'] - 407.00) < 0.005, res['gross']
    # 退款 = 82.90（全额）+ 20.00（实际部分退款） = 102.90
    assert abs(res['refund'] - 102.90) < 0.005, res['refund']
    # 净支出 = 407.00 - 102.90 = 304.10
    assert abs(res['net'] - 304.10) < 0.005, res['net']
    assert res['counts']['refunded'] == 2


def test_partial_refund_state_detected():
    """'已退款¥20.00' 这种部分退款也要被识别为已退款。"""
    rows, _, _ = _load()
    partial = [r for r in rows if '¥' in str(r['当前状态'])]
    assert len(partial) == 1
    assert partial[0]['已退款'] is True


# ---------------------------------------------------------------------------
# 测试 3：分类完整性（问题 2 的根因防线）
# ---------------------------------------------------------------------------
def test_category_sums_to_net():
    """分类金额必须合计等于净支出，笔数必须等于净支出笔数。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)
    cat_sum = round(sum(v['金额'] for v in res['categories'].values()), 2)
    cat_cnt = sum(v['笔数'] for v in res['categories'].values())
    assert abs(cat_sum - res['net']) < 0.005, f'分类合计 {cat_sum} != 净支出 {res["net"]}'
    assert cat_cnt == res['counts']['net_expense'], f'{cat_cnt} != {res["counts"]["net_expense"]}'


def test_qr_payment_reclassified_by_payee():
    """扫码付款给羽毛球馆应归入运动健身，而不是转账与红包。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)
    assert '⑥ 运动健身' in res['categories'], res['categories'].keys()
    assert abs(res['categories']['⑥ 运动健身']['金额'] - 6.00) < 0.005
    # 真正的转账仍在转账类
    assert abs(res['categories']['⑦ 转账与红包']['金额'] - 30.00) < 0.005


# ---------------------------------------------------------------------------
# 测试 4：收入拆分与亲人识别
# ---------------------------------------------------------------------------
def test_family_transfer_separated():
    """--family 指定的对方应单独归为"亲人/资助转入"。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows, family=['母亲'])
    assert '亲人/资助转入' in res['income_breakdown'], res['income_breakdown']
    assert abs(res['income_breakdown']['亲人/资助转入']['金额'] - 1000.00) < 0.005
    # 不加 family 时不应出现该类别
    res2 = analyze.summarize(rows)
    assert '亲人/资助转入' not in res2['income_breakdown']


def test_slash_rows_excluded_from_both():
    """收/支 为 '/' 的行（提现/充值）既不算支出也不算收入。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)
    assert res['counts']['other'] == 1
    assert res['counts']['expense'] == 8
    assert res['counts']['income'] == 1


# ---------------------------------------------------------------------------
# 测试 5：预付充值提示
# ---------------------------------------------------------------------------
def test_prepaid_detected():
    """校园卡充值属于预付性支出，应能被识别出来单独提示。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)
    pre = [r for r in res['net_rows']
           if any(k in str(r['商品']) + str(r['交易对方']) for k in ('充值', '开卡', '校园卡'))]
    assert len(pre) == 1, pre
    assert abs(pre[0]['金额'] - 200.00) < 0.005


# ---------------------------------------------------------------------------
# 测试 6：分段统计的完整性断言（问题 2 同类）
# ---------------------------------------------------------------------------
def test_segments_cover_all_expenses():
    """分段必须完整覆盖净支出；分段合计要等于总净支出。"""
    rows, _, _ = _load()
    segs = [
        {'name': 'A 假期', 'lo': '2026-09-03', 'hi': '2026-09-05', 'kind': '假期'},
        {'name': 'B 在校', 'lo': '2026-09-05', 'hi': '2026-09-11', 'kind': '在校'},
    ]
    res = timeseg.segment(rows, segs)
    assert not res['missed'], f'存在未覆盖记录：{len(res["missed"])} 笔'
    seg_sum = round(sum(s['amt'] for s in res['segments']), 2)
    assert abs(seg_sum - res['total']) < 0.005, f'{seg_sum} != {res["total"]}'
    assert res['by_kind']['假期']['amt'] + res['by_kind']['在校']['amt'] == res['total']


def test_segments_report_gaps():
    """分段有空隙时，未覆盖的记录必须被检出（而不是静默丢失）。"""
    rows, _, _ = _load()
    segs = [{'name': '只覆盖一部分', 'lo': '2026-09-05', 'hi': '2026-09-06', 'kind': '在校'}]
    res = timeseg.segment(rows, segs)
    assert res['missed'], '有空隙时应检出未覆盖记录'
    covered = round(sum(s['amt'] for s in res['segments']), 2)
    missed = round(sum(r['金额'] for r in res['missed']), 2)
    assert abs(covered + missed - res['total']) < 0.005, '已覆盖 + 未覆盖 应等于总额'


# ---------------------------------------------------------------------------
# 测试 7：餐饮识别与预算恒等式（问题 5）
# ---------------------------------------------------------------------------
def test_meals_excludes_non_food():
    """药店等含"药/食"字样的支出不能被算作餐饮。"""
    rows, _, _ = _load()
    m = meals.analyze_meals(rows)
    # 夹具里没有药店；转而验证食堂被正确识别
    assert '食堂' in m['by_channel'], m['by_channel']
    assert m['anchor']['早餐'][1] >= 1, '应识别到食堂早餐'


def test_budget_food_plus_other_equals_recurring():
    """恒等式：餐饮 + 非餐饮 = 经常性支出（防止同一笔钱被算两遍）。"""
    rows, _, _ = _load()
    b = budget.analyze_budget(rows, month_days=31)
    lhs = round(b['food_recurring'] + b['other_recurring'], 2)
    assert abs(lhs - b['recurring']) < 0.005, \
        f'餐饮 {b["food_recurring"]} + 非餐饮 {b["other_recurring"]} != 经常性 {b["recurring"]}'


def test_budget_recurring_plus_oneoff_equals_total():
    """恒等式：经常性 + 一次性 = 窗口净支出。"""
    rows, _, _ = _load()
    b = budget.analyze_budget(rows, month_days=31)
    lhs = round(b['recurring'] + b['one_off'], 2)
    assert abs(lhs - b['total']) < 0.005, f'{lhs} != {b["total"]}'


def test_budget_cuts_not_empty():
    """有可压缩类别时，省钱方案不能是空的（问题 4 的回归测试）。"""
    rows, _, _ = _load()
    b = budget.analyze_budget(rows, month_days=31)
    # 夹具含网约车 25.60 元，属于可压缩项
    assert b['cuts'], '存在网约车支出，省钱方案不应为空'
    assert b['total_save'] > 0, b['total_save']
    keys = [c['键'] for c in b['cuts']]
    assert '网约车' in keys, keys


# ---------------------------------------------------------------------------
# 测试 8：Excel 报告中占比列必须合计 100%（问题 2 的直接回归测试）
# ---------------------------------------------------------------------------
def _eval_cell(ws, coord, cache=None):
    """对 xlsx 里的单元格做真求值（支持 SUM 与四则运算），而不是只看公式字符串。

    这是本测试文件的关键工具：早期版本只断言"公式长什么样"，
    结果 timeseg.py 里分母指向错误行的缺陷从测试里溜了过去。
    """
    if cache is None:
        cache = {}
    if coord in cache:
        return cache[coord]
    v = ws[coord].value
    if isinstance(v, (int, float)):
        cache[coord] = v
        return v
    if not (isinstance(v, str) and v.startswith('=')):
        cache[coord] = v
        return v
    body = v[1:]
    cache[coord] = 0  # 防循环
    m = re.fullmatch(r'SUM\(([A-Z]+\d+):([A-Z]+\d+)\)', body)
    if m:
        c0, r0 = re.match(r'([A-Z]+)(\d+)', m.group(1)).groups()
        c1, r1 = re.match(r'([A-Z]+)(\d+)', m.group(2)).groups()
        c0i, c1i = ord(c0) - 65, ord(c1) - 65
        tot = 0.0
        for rr in range(int(r0), int(r1) + 1):
            for cc in range(c0i, c1i + 1):
                x = _eval_cell(ws, f'{chr(65 + cc)}{rr}', cache)
                if isinstance(x, (int, float)):
                    tot += x
        cache[coord] = tot
        return tot
    expr = re.sub(r'\$?([A-Z]{1,3}\$?\d+)',
                  lambda mm: str(_eval_cell(ws, mm.group(1).replace('$', ''), cache)), body)
    try:
        result = eval(expr)  # 仅用于测试内的公式求值
    except Exception:
        result = None
    cache[coord] = result
    return result


def _find_table(ws, header_text):
    """定位一张表：返回 (表头行, 数据首行, 合计行)。"""
    hdr = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == header_text:
            hdr = r
            break
    assert hdr, f'未找到表头 {header_text!r}'
    first = hdr + 1
    last = first
    while ws.cell(row=last, column=1).value not in (None, '合计'):
        last += 1
    return hdr, first, last


def test_report_percentage_column_sums_to_100():
    """生成报告后**逐格求值**：分类占比列合计必须为 100%（不是只看公式形状）。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    try:
        analyze.write_report(res, path, title='测试报告')
        wb = load_workbook(path)
        assert '统计概览' in wb.sheetnames, wb.sheetnames
        ws = wb['统计概览']
        hdr, first, total_row = _find_table(ws, '类别')

        # 公式形状：分母必须指向合计行
        for r in range(first, total_row):
            f = ws.cell(row=r, column=3).value
            assert isinstance(f, str) and f.startswith('='), f'C{r} 不是公式: {f!r}'
            assert f == f'=B{r}/$B${total_row}', f'C{r} 公式不正确: {f}'
        assert ws.cell(row=total_row, column=3).value == '=1'
        assert ws.cell(row=total_row, column=4).value == \
            f'=SUM(D{first}:D{total_row - 1})'

        # ★ 真求值：占比列合计必须为 100%
        cache = {}
        s = sum(_eval_cell(ws, f'C{r}', cache) for r in range(first, total_row))
        assert abs(s - 1) < 1e-9, f'占比列合计 = {s * 100:.2f}%，应为 100%'
        # 金额列合计必须等于净支出
        amt = _eval_cell(ws, f'B{total_row}', cache)
        assert abs(amt - res['net']) < 0.005, f'金额合计 {amt} != 净支出 {res["net"]}'
    finally:
        os.remove(path)


def test_segment_report_percentage_column_sums_to_100():
    """分段报告的占比列也必须真求值到 100%（timeseg.py 曾漏掉 +1 导致整列全错）。"""
    rows, _, _ = _load()
    segs = [
        {'name': 'A 假期', 'lo': '2026-09-03', 'hi': '2026-09-05', 'kind': '假期'},
        {'name': 'B 在校', 'lo': '2026-09-05', 'hi': '2026-09-11', 'kind': '在校'},
    ]
    res = timeseg.segment(rows, segs)
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    try:
        timeseg.write_seg_report(res, path)
        ws = load_workbook(path)['分段统计']
        hdr, first, total_row = _find_table(ws, '时段')
        nseg = total_row - first
        assert nseg == len(segs)
        # 公式形状：分母必须是合计行（= 首行 + 时段数）
        expect_total = first + nseg
        assert total_row == expect_total, f'合计行 {total_row} != {expect_total}'
        for r in range(first, total_row):
            f = ws.cell(row=r, column=7).value
            assert f == f'=E{r}/$E${total_row}', \
                f'G{r} 公式不正确（分母应指向合计行 {total_row}）: {f}'
        # ★ 真求值
        cache = {}
        s = sum(_eval_cell(ws, f'G{r}', cache) for r in range(first, total_row))
        assert abs(s - 1) < 1e-9, f'分段占比列合计 = {s * 100:.2f}%，应为 100%'
    finally:
        os.remove(path)


def test_balance_uses_net_not_gross():
    """结余必须是"收入 − 净支出"，不能减含退款的流水（否则退款被算两遍）。"""
    rows, _, _ = _load()
    res = analyze.summarize(rows)
    expect = round(res['income'] - res['net'], 2)
    assert abs(res['balance'] - expect) < 0.005, \
        f"结余 {res['balance']} 应等于 收入 {res['income']} − 净支出 {res['net']} = {expect}"
    # 关键反例：若错误地减 gross，会正好差一个退款额
    wrong = round(res['income'] - res['gross'], 2)
    assert res['refund'] > 0, '夹具应包含退款，否则这条测试没有区分力'
    assert abs(res['balance'] - wrong) > 0.005, '结余不能等于 收入 − 流水'


def test_income_only_bill_does_not_crash():
    """只有收入、没有支出的账单：预算与餐饮分析不能崩溃。"""
    from datetime import datetime
    rows = [
        (datetime(2026, 9, 5, 12, 0), '转账', '母亲', '转账备注:微信转账',
         '收入', 1500.00, '/', '已存入零钱'),
    ]
    r, _, _ = _load_custom(rows)
    # analyze 本身
    res = analyze.summarize(r)
    assert res['net'] == 0 and res['income'] == 1500.00
    # 这三处曾因 min() 空序列抛 ValueError
    b = budget.analyze_budget(r)
    assert b['recurring'] == 0 and b['cuts'] == []
    m = meals.analyze_meals(r)
    assert m['total'] == 0
    assert m['anchor']['早餐'][0] == 0
    base, plans, unit = meals.budget_plans(m)
    assert base == 0 and len(plans) > 0


# ---------------------------------------------------------------------------
# 测试 9：自定义规则文件的健壮性
# ---------------------------------------------------------------------------
def _write_tmp(text, encoding='utf-8'):
    fd, path = tempfile.mkstemp(suffix='.json')
    os.close(fd)
    with open(path, 'w', encoding=encoding) as f:
        f.write(text)
    return path


def test_rules_with_bom_and_leading_underscore_notes():
    """规则文件带 BOM（Windows PowerShell 默认写 BOM）要能读；
    以下划线开头的说明键要被忽略，不能当成一个类别。"""
    body = ('{'
            '"_说明": ["这是说明，不应成为类别"],'
            '"① 餐饮美食": ["示例大学"],'
            '"② 交通出行": ["如祺出行"]'
            '}')
    path = _write_tmp(body, encoding='utf-8-sig')   # 带 BOM
    try:
        rules = analyze.load_rules(path)
        cats = [c for c, _ in rules]
        assert '① 餐饮美食' in cats and '② 交通出行' in cats, cats
        assert '_说明' not in cats, f'说明键不应成为类别: {cats}'
        # 且这些规则真的能用于分类
        rows, _, _ = _load()
        res = analyze.summarize(rows, rules=rules)
        assert '① 餐饮美食' in res['categories']
        assert abs(sum(v['金额'] for v in res['categories'].values()) - res['net']) < 0.005
    finally:
        os.remove(path)


def test_rules_invalid_json_raises_clear_error():
    """非法 JSON 应给出带格式示例的明确错误，而不是抛原始 JSONDecodeError。"""
    path = _write_tmp('{ this is not json')
    try:
        try:
            analyze.load_rules(path)
        except SystemExit as e:
            msg = str(e)
            assert '不是合法 JSON' in msg, msg
            assert '餐饮美食' in msg, '错误信息里应给出格式示例'
            return
        raise AssertionError('非法 JSON 应抛出 SystemExit')
    finally:
        os.remove(path)


def test_rules_wrong_value_type_raises():
    """值不是字符串数组时应明确报错，而不是静默把规则文件当成有效。"""
    path = _write_tmp('{"① 餐饮美食": "不是数组"}')
    try:
        try:
            analyze.load_rules(path)
        except SystemExit as e:
            assert '必须是字符串数组' in str(e), str(e)
            return
        raise AssertionError('值类型错误时应抛出 SystemExit')
    finally:
        os.remove(path)


def test_rules_only_notes_raises():
    """只有说明键（没有真实类别）时应报错，而不是产出空规则集。"""
    path = _write_tmp('{"_说明": ["x"]}')
    try:
        try:
            analyze.load_rules(path)
        except SystemExit as e:
            assert '没有任何有效类别' in str(e), str(e)
            return
        raise AssertionError('空规则集时应抛出 SystemExit')
    finally:
        os.remove(path)


def test_shipped_rules_example_is_valid():
    """仓库里随附的示例规则文件必须能直接被 load_rules 读取。"""
    p = os.path.join(ROOT, 'examples', 'rules.example.json')
    if not os.path.exists(p):
        raise AssertionError(f'缺少示例规则文件：{p}')
    rules = analyze.load_rules(p)
    assert len(rules) >= 8, f'示例规则类别过少: {len(rules)}'
    for cat, kws in rules:
        assert not cat.startswith('_'), f'示例文件不应把说明键当类别: {cat}'
        assert kws, f'类别 {cat} 的关键词为空'


# ---------------------------------------------------------------------------
# 测试 10：餐饮口径唯一性（问题 5 的回归防线）
# ---------------------------------------------------------------------------
def test_food_classification_is_single_source_of_truth():
    """meals 与 budget 对"什么算餐饮"必须给出完全一致的答案。

    历史上两者各有一份判断，对同一批数据差了 572.60 元：
    budget 额外排除了"类别以 ⑫ 开头"的记录。现在两者都只用 analyze.chow_kind。
    """
    rows, _, _ = _load()
    rules = analyze.compile_rules()
    for r in rows:
        r.setdefault('类别', analyze.classify(r, rules))

    m = meals.analyze_meals(rows)
    b = budget.analyze_budget(rows)

    meals_food = round(sum(v['金额'] for v in m['by_channel'].values()), 2)
    # budget 会在非一次性支出上统计渠道；夹具里餐饮都不属于一次性项
    assert abs(b['food_recurring'] - meals_food) < 0.005, (
        f"餐饮口径不一致：meals {meals_food} vs budget {b['food_recurring']}")

    # 逐笔核对：两边认定的餐饮记录集合应完全相同
    meal_rows = {(r['iso'],r['交易对方'],r['金额']) for r in m['food']}
    budget_food = {(r['iso'],r['交易对方'],r['金额']) for r in analyze.net_expenses(rows) if analyze.chow_kind(r)}
    assert meal_rows == budget_food, (
        f'餐饮记录集合不同：仅 meals 有 {len(meal_rows - budget_food)} 笔，'
        f'仅 budget 有 {len(budget_food - meal_rows)} 笔')


def test_food_module_reuses_analyze_implementation():
    """meals 与 budget 不应各自定义 chow_kind，必须复用 analyze 里的那一份。"""
    assert meals.chow_kind is analyze.chow_kind, \
        'meals.chow_kind 必须是 analyze.chow_kind，否则又会出现两套口径'
    # budget 内部通过 analyze 导入，因此同一批数据结果一致
    rows, _, _ = _load()
    rules = analyze.compile_rules()
    for r in rows:
        r.setdefault('类别', analyze.classify(r, rules))
    sample = rows[0]
    assert meals.chow_kind(sample) == analyze.chow_kind(sample)


def test_food_plus_nonfood_equals_recurring_on_synthetic_data():
    """恒等式在演示数据上也必须成立：餐饮 + 非餐饮 = 经常性支出。"""
    p = os.path.join(ROOT, 'examples', 'data', 'demo_bill.xlsx')
    if not os.path.exists(p):
        raise AssertionError(f'缺少演示数据：{p}')
    rows, _, _ = analyze.load_bill(p)
    b = budget.analyze_budget(rows)
    lhs = round(b['food_recurring'] + b['other_recurring'], 2)
    assert abs(lhs - b['recurring']) < 0.005, \
        f"餐饮 {b['food_recurring']} + 非餐饮 {b['other_recurring']} != 经常性 {b['recurring']}"





def _make_custom_bill(path, rows, preamble_lines=13):
    """用任意数据行写一个账单（供边界测试使用）。"""
    wb = Workbook()
    ws = wb.active
    for i in range(1, preamble_lines + 1):
        ws.cell(row=i, column=1, value=f'账单抬头第 {i} 行')
    hr = preamble_lines + 1
    for j, h in enumerate(HEADERS, 1):
        ws.cell(row=hr, column=j, value=h)
    for i, r in enumerate(rows, hr + 1):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
    wb.save(path)


def _load_custom(rows, preamble_lines=13):
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    try:
        _make_custom_bill(path, rows, preamble_lines)
        return analyze.load_bill(path)
    finally:
        os.remove(path)


# ---------------------------------------------------------------------------
# 测试 8：边界情况（空数据、全收入、损坏表头）
# ---------------------------------------------------------------------------
def test_empty_bill_does_not_crash():
    """完全没有数据行的账单不能崩溃，且各合计为 0。"""
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    try:
        _make_custom_bill(path, [])
        rows, headers, hr = analyze.load_bill(path)
        assert rows == [], f'应读到 0 条，实际 {len(rows)}'
        res = analyze.summarize(rows)
        assert res['net'] == 0 and res['gross'] == 0 and res['income'] == 0
        assert res['categories'] == {}
        assert res['daily'] == 0, '空账单的日均应为 0，不能除零'
    finally:
        os.remove(path)


def test_income_only_bill():
    """只有收入的账单：净支出为 0，收入与结余正确。"""
    from datetime import datetime
    rows = [
        (datetime(2026, 9, 5, 12, 0), '转账', '母亲', '转账备注:微信转账',
         '收入', 1500.00, '/', '已存入零钱'),
        (datetime(2026, 9, 6, 12, 0), '微信红包', '同学', '/',
         '收入', 66.00, '/', '已存入零钱'),
    ]
    r, _, _ = _load_custom(rows)
    res = analyze.summarize(r)
    assert res['net'] == 0, res['net']
    assert abs(res['income'] - 1566.00) < 0.005, res['income']
    assert res['categories'] == {}
    # 收入构成应按类型拆分：普通转账归"其他转入"，红包单独一类
    assert '其他转入' in res['income_breakdown'], res['income_breakdown']
    assert '红包收入' in res['income_breakdown'], res['income_breakdown']
    assert abs(res['income_breakdown']['其他转入']['金额'] - 1500.00) < 0.005
    assert abs(res['income_breakdown']['红包收入']['金额'] - 66.00) < 0.005
    # 拆分后的金额必须等于收入总额（防止漏分类）
    assert abs(sum(v['金额'] for v in res['income_breakdown'].values())
               - res['income']) < 0.005


def test_all_refunded_bill():
    """所有支出都已退款：流水有金额，但净支出为 0。"""
    from datetime import datetime
    rows = [
        (datetime(2026, 9, 5, 12, 0), '商户消费', '某店', '商品',
         '支出', 50.00, '零钱', '已全额退款'),
        (datetime(2026, 9, 6, 12, 0), '商户消费', '某店', '商品',
         '支出', 30.00, '零钱', '已全额退款'),
    ]
    r, _, _ = _load_custom(rows)
    res = analyze.summarize(r)
    assert abs(res['gross'] - 80.00) < 0.005, res['gross']
    assert abs(res['refund'] - 80.00) < 0.005, res['refund']
    assert res['net'] == 0, res['net']
    assert res['counts']['refunded'] == 2
    assert res['categories'] == {}


def test_missing_header_raises_clear_error():
    """没有表头的文件应抛出明确错误，而不是静默返回空结果。"""
    wb = Workbook()
    ws = wb.active
    for i in range(1, 6):
        ws.cell(row=i, column=1, value=f'这不是账单第 {i} 行')
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    try:
        wb.save(path)
        try:
            analyze.load_bill(path)
        except ValueError as e:
            assert '表头' in str(e), f'错误信息应说明是表头问题，实际：{e}'
            return
        raise AssertionError('缺少表头时应抛出 ValueError')
    finally:
        os.remove(path)


def test_amount_field_types():
    """金额应被规范为 float 且保留两位小数，日期应被解析为 datetime。"""
    import datetime as dt
    rows, _, _ = _load()
    assert len(rows) == 10
    for r in rows:
        assert isinstance(r['金额'], float), f"金额不是 float: {r['金额']!r}"
        assert round(r['金额'], 2) == r['金额'], f"金额未保留两位小数: {r['金额']}"
        assert isinstance(r['dt'], dt.datetime), f"dt 不是 datetime: {r['dt']!r}"
        assert isinstance(r['iso'], str) and len(r['iso']) == 10
        assert isinstance(r['已退款'], bool)


def test_unassigned_expense_detected_in_segments():
    """分段完全没覆盖任何记录时，未覆盖金额应等于总额（不能静默为 0）。"""
    rows, _, _ = _load()
    segs = [{'name': '完全错位的时段', 'lo': '2020-01-01', 'hi': '2020-01-02', 'kind': '假期'}]
    res = timeseg.segment(rows, segs)
    assert res['segments'][0]['amt'] == 0
    missed_amt = round(sum(r['金额'] for r in res['missed']), 2)
    assert abs(missed_amt - res['total']) < 0.005, \
        f'未覆盖金额 {missed_amt} 应等于总额 {res["total"]}'



# ---------------------------------------------------------------------------
# 内置极简 runner（不依赖 pytest）
# ---------------------------------------------------------------------------

def _case(amount=39, status='支付成功', direction='支出', typ='商户消费', party='食堂'):
    from datetime import datetime
    t=datetime(2026,9,15,12)
    return {'dt':t,'交易时间':t,'iso':'2026-09-15','hour':12,'金额':float(amount),
            '当前状态':status,'已退款':'退款' in status,'收/支':direction,
            '交易类型':typ,'交易对方':party,'商品':'午餐','支付方式':'零钱'}

def test_actual_partial_refund_amount():
    r=analyze.summarize([_case(status='已退款(¥20.00)')])
    assert (r['gross'],r['refund'],r['net']) == (39,20,19)
    assert meals.analyze_meals([_case(status='已退款(¥20.00)')])['total']==19
    assert budget.analyze_budget([_case(status='已退款(¥20.00)')])['total']==19

def test_refund_income_is_not_earned_twice():
    r=analyze.summarize([_case(100,'已全额退款'),_case(100,'已存入零钱','收入','商户-退款')])
    assert r['balance']==0 and r['cash_balance']==0 and r['income']==0

def test_unknown_refund_rejected():
    try: analyze.summarize([_case(status='退款处理中')])
    except ValueError: return
    raise AssertionError('含糊退款状态不得猜测金额')

def test_explicit_period_includes_zero_spend_days():
    rows=[_case(30)]
    assert analyze.summarize(rows,lo='2026-09-01',hi='2026-10-01')['daily']==1
    assert meals.analyze_meals(rows,lo='2026-09-01',hi='2026-10-01')['daily']==1
    assert budget.analyze_budget(rows,lo='2026-09-01',hi='2026-10-01')['days']==30

def test_text_dates_parsed_and_bad_dates_rejected():
    row=('2026-09-15 12:00:00','商户消费','食堂','午餐','支出',30,'零钱','支付成功')
    rows,_,_=_load_custom([row]); assert len(rows)==1
    try: _load_custom([('bad-date',)+row[1:]])
    except ValueError: return
    raise AssertionError('无效日期必须报错')

def _prices():
    return {'anchor':{'早餐':(5,1),'午餐':(10,1),'晚餐':(10,1)},
            'unit':{'奶茶咖啡':(8,1),'加餐夜宵':(4,1),'外卖正餐':(20,1)}}

def test_takeout_replaces_canteen_lunch():
    p=meals.budget_plans(_prices(),7)[1][1]
    assert p['base']==165 and p['out']==20 and p['total']==201

def test_missing_prices_are_unavailable():
    p=meals.budget_plans(meals.analyze_meals([]),7)[1]
    assert all(not x['available'] and x['missing'] and x['total'] is None for x in p)

def test_empty_cli_paths():
    import contextlib,io
    with contextlib.redirect_stdout(io.StringIO()):
        m=meals.analyze_meals([]); base,plans,_=meals.budget_plans(m)
        meals.print_meals(m,base,plans)
        budget.print_budget(budget.analyze_budget([]),m,base,plans)
        timeseg.print_seg(timeseg.segment([],[{'name':'x','lo':'2026-09-01','hi':'2026-10-01'}]))

def test_budget_food_savings_not_deducted_from_other():
    import contextlib,io
    b=budget.analyze_budget([_case(20,party='奶茶')]); base,plans,_=meals.budget_plans(_prices(),7)
    out=io.StringIO()
    with contextlib.redirect_stdout(out): budget.print_budget(b,None,base,plans)
    assert '非餐饮月化 0.00 − 非餐饮节省 0.00 = 0.00' in out.getvalue()

def test_segment_overlap_rejected_even_without_transactions():
    try: timeseg.segment([],[{'lo':'2026-09-01','hi':'2026-09-10'},{'lo':'2026-09-05','hi':'2026-09-12'}])
    except ValueError: return
    raise AssertionError('空账单也要验证分段')

def test_nonpositive_budget_days_rejected():
    for n in (0,-1):
        try: meals.budget_plans(_prices(),n)
        except ValueError: continue
        raise AssertionError('预算天数应为正数')

def test_empty_excel_totals_have_no_circular_reference():
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        p=os.path.join(d,'empty.xlsx');analyze.write_report(analyze.summarize([]),p)
        w=load_workbook(p);s=w['统计概览']
        for row in s:
            if row[0].value=='合计':
                assert row[1].value=='=0' and row[2].value=='=0'


def _run_all():
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith('test_') and callable(f)]
    passed, failed = 0, []
    print(f'运行 {len(tests)} 个测试\n' + '=' * 68)
    for name, fn in tests:
        try:
            fn()
            passed += 1
            print(f'  PASS  {name}')
        except AssertionError as e:
            failed.append((name, f'断言失败: {e}'))
            print(f'  FAIL  {name}\n        {e}')
        except Exception as e:                       # noqa: BLE001
            failed.append((name, f'{type(e).__name__}: {e}'))
            print(f'  ERROR {name}\n        {type(e).__name__}: {e}')
    print('=' * 68)
    print(f'通过 {passed}/{len(tests)}')
    if failed:
        print('\n失败项：')
        for n, why in failed:
            print(f'  - {n}: {why}')
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(_run_all())
